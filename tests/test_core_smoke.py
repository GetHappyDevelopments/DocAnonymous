from pathlib import Path
import shutil
import zipfile

from doc_anonymizer.app.core.anonymizer import DocumentAnonymizer
from doc_anonymizer.app.core.models import DocumentJob, Finding
from doc_anonymizer.app.core.restorer import DocumentRestorer
from doc_anonymizer.app.core.scanner import DocumentScanner
from doc_anonymizer.app.formats import get_handler
from doc_anonymizer.app.storage.project_state import ProjectStateStore


def _write_minimal_text_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({escaped}) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream".encode("latin-1"),
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode("ascii"))
        data.extend(obj)
        data.extend(b"\nendobj\n")
    xref_offset = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(bytes(data))


def test_txt_scan_and_anonymize(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text(
        "Bayer AG, Vater&Soehne GmbH, Herr Max Mustermann, test@example.com",
        encoding="utf-8",
    )
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    originals = {finding.original_text for finding in job.findings}

    assert "Bayer AG" in originals
    assert "Vater&Soehne GmbH" in originals
    assert "Herr Max Mustermann" in originals
    assert "test@example.com" in originals

    DocumentAnonymizer().anonymize(job, tmp_path)

    text = job.output_path.read_text(encoding="utf-8")
    assert "Firma AG 1" in text
    assert "Firma GmbH 1" in text
    assert "Person 1" in text
    assert "E-Mail 1" in text
    assert job.restore_package_path.exists()
    assert job.restore_package_path.suffix == ".dam"
    assert zipfile.is_zipfile(job.restore_package_path)


def test_company_variant_and_authority_detection(tmp_path: Path) -> None:
    source = tmp_path / "variants.txt"
    source.write_text(
        "Bayer Aktiengesellschaft arbeitet mit Bayer AG und Amtsgericht Muenchen.",
        encoding="utf-8",
    )
    job = DocumentJob(source)

    DocumentScanner().scan(job)

    company_replacements = {
        finding.replacement_text
        for finding in job.findings
        if finding.category == "company"
    }
    assert company_replacements == {"Firma AG 1"}
    assert any(finding.category == "authority" for finding in job.findings)


def test_umlaut_person_company_and_address_detection(tmp_path: Path) -> None:
    source = tmp_path / "umlauts.txt"
    source.write_text(
        "Frau Dr. Anna Müller wohnt in der Käferstraße 12, 12345 München.\n"
        "Die Müller & Söhne GmbH arbeitet später nur als Müller & Söhne weiter.",
        encoding="utf-8",
    )
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    originals = {finding.original_text for finding in job.findings}

    assert "Frau Dr. Anna Müller" in originals
    assert "Käferstraße 12, 12345 München" in originals
    assert "Müller & Söhne GmbH" in originals
    assert "Müller & Söhne" in originals


def test_multiline_address_block_detection(tmp_path: Path) -> None:
    source = tmp_path / "address.txt"
    source.write_text(
        "Max Mustermann\nMusterweg 7\n10115 Berlin\n\nBitte zeitnah melden.",
        encoding="utf-8",
    )
    job = DocumentJob(source)

    DocumentScanner().scan(job)

    assert any(
        finding.category == "address"
        and "Max Mustermann" in finding.original_text
        and "10115 Berlin" in finding.original_text
        for finding in job.findings
    )


def test_scan_does_not_merge_person_name_across_control_character(tmp_path: Path) -> None:
    source = tmp_path / "control.txt"
    source.write_text("Stefan Jobst\x0bLeiter", encoding="utf-8")
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    originals = {finding.original_text for finding in job.findings}

    assert "Stefan Jobst" in originals
    assert "Stefan Jobst Leiter" not in originals
    assert not any("Leiter" in original for original in originals)
    assert all(
        not any(ord(char) < 32 and char not in "\r\n\t" for char in finding.original_text)
        for finding in job.findings
    )


def test_manual_findings_are_reused_as_local_learning(tmp_path: Path) -> None:
    source = tmp_path / "learned.txt"
    source.write_text("Codexia Projektbüro liefert. Codexia Projektbüro rechnet ab.", encoding="utf-8")
    job = DocumentJob(source)
    job.findings.append(
        DocumentScanner()._merge_and_assign_placeholders(
            [
                Finding(
                    original_text="Codexia Projektbüro",
                    replacement_text="",
                    category="company",
                    confidence=1.0,
                    source="manual",
                )
            ]
        )[0]
    )

    DocumentScanner().scan(job)

    learned = [finding for finding in job.findings if finding.original_text == "Codexia Projektbüro"]
    assert learned
    assert learned[0].occurrence_count == 2


def test_project_state_roundtrip(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text("Bayer AG", encoding="utf-8")
    job = DocumentJob(source)
    DocumentScanner().scan(job)
    job.findings[0].correct = True
    job.findings[0].incorrect = False
    project = tmp_path / "state.docanon"

    store = ProjectStateStore()
    store.save(project, [job])
    loaded = store.load(project)

    assert loaded[0].source_path == source
    assert loaded[0].findings[0].replacement_text == "Firma AG 1"
    assert loaded[0].findings[0].correct is True
    assert loaded[0].findings[0].incorrect is False


def test_project_state_load_disables_incorrect_findings(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text("Bayer AG", encoding="utf-8")
    project = tmp_path / "state.docanon"
    project.write_text(
        """{
  "schemaVersion": "1.0",
  "jobs": [
    {
      "source_path": "%s",
      "status": "Scan abgeschlossen",
      "findings": [
        {
          "original_text": "Bayer AG",
          "replacement_text": "Firma AG 1",
          "category": "company",
          "enabled": true,
          "correct": true,
          "incorrect": true
        }
      ]
    }
  ]
}"""
        % str(source).replace("\\", "\\\\"),
        encoding="utf-8",
    )

    loaded = ProjectStateStore().load(project)

    assert loaded[0].findings[0].incorrect is True
    assert loaded[0].findings[0].correct is False
    assert loaded[0].findings[0].enabled is False


def test_rescan_preserves_incorrect_findings_as_inactive_at_end(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text("Bayer AG test@example.com", encoding="utf-8")
    job = DocumentJob(source)
    scanner = DocumentScanner()
    scanner.scan(job)

    bayer = next(finding for finding in job.findings if finding.original_text == "Bayer AG")
    bayer.incorrect = True
    bayer.enabled = False

    scanner.scan(job)

    assert job.findings[-1].original_text == "Bayer AG"
    assert job.findings[-1].incorrect is True
    assert job.findings[-1].enabled is False


def test_txt_restore_roundtrip(tmp_path: Path) -> None:
    source = tmp_path / "restore.txt"
    source.write_text("Bayer AG und Herr Max Mustermann", encoding="utf-8")
    job = DocumentJob(source)
    DocumentScanner().scan(job)
    DocumentAnonymizer().anonymize(job, tmp_path)

    restored = DocumentRestorer().restore(job.output_path, job.restore_package_path, tmp_path / "restored.txt")

    assert restored.read_text(encoding="utf-8") == source.read_text(encoding="utf-8")


def test_docx_scan_anonymize_and_restore(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Bayer AG")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Vater&Soehne GmbH"
    doc.save(source)
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    DocumentAnonymizer().anonymize(job, tmp_path)
    anonymized_text = "\n".join(chunk.text for chunk in get_handler(job.output_path).extract_text(job.output_path))

    assert "Firma AG 1" in anonymized_text
    assert "Firma GmbH 1" in anonymized_text

    restored = DocumentRestorer().restore(job.output_path, job.restore_package_path, tmp_path / "restored.docx")
    restored_text = "\n".join(chunk.text for chunk in get_handler(restored).extract_text(restored))
    assert "Bayer AG" in restored_text
    assert "Vater&Soehne GmbH" in restored_text


def test_docx_anonymize_removes_header_images(tmp_path: Path) -> None:
    from docx import Document
    from PIL import Image

    source = tmp_path / "header-image.docx"
    logo = tmp_path / "logo.png"
    Image.new("RGB", (32, 16), (120, 120, 120)).save(logo)

    doc = Document()
    doc.add_paragraph("Bayer AG")
    doc.sections[0].header.paragraphs[0].add_run().add_picture(str(logo))
    doc.save(source)
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        names = zf.namelist()
        xml_parts = {
            name: zf.read(name)
            for name in names
            if name.endswith((".xml", ".rels")) and name.startswith("word/")
        }

    assert not any("/media/" in name.replace("\\", "/") for name in names)
    assert not any(b"/relationships/image" in data for data in xml_parts.values())
    assert not any(b"<a:blip" in data or b"<v:imagedata" in data for data in xml_parts.values())
    assert not any(b"<w:drawing" in data for data in xml_parts.values())
    assert job.image_replacements


def test_docx_anonymize_can_keep_selected_image(tmp_path: Path) -> None:
    from docx import Document
    from PIL import Image

    from doc_anonymizer.app.formats.common import collect_openxml_images

    source = tmp_path / "selected-image.docx"
    keep_logo = tmp_path / "keep.png"
    remove_logo = tmp_path / "remove.png"
    Image.new("RGB", (32, 16), (120, 120, 120)).save(keep_logo)
    Image.new("RGB", (32, 16), (200, 200, 200)).save(remove_logo)

    doc = Document()
    doc.add_paragraph("Bayer AG")
    doc.add_picture(str(keep_logo))
    doc.add_picture(str(remove_logo))
    doc.save(source)

    job = DocumentJob(source)
    DocumentScanner().scan(job)
    job.image_replacements = collect_openxml_images(source)
    job.image_replacements[0].keep = True
    kept_path = job.image_replacements[0].locations[0].extra["package_path"]
    removed_path = job.image_replacements[1].locations[0].extra["package_path"]

    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        names = zf.namelist()

    assert kept_path in names
    assert removed_path not in names
    assert [image.keep for image in job.image_replacements] == [True, False]


def test_docx_restore_rebuilds_removed_image_parts(tmp_path: Path) -> None:
    from docx import Document
    from PIL import Image

    source = tmp_path / "restore-image.docx"
    logo = tmp_path / "logo.png"
    Image.new("RGB", (32, 16), (120, 120, 120)).save(logo)

    doc = Document()
    doc.add_paragraph("Bayer AG")
    doc.add_picture(str(logo))
    doc.save(source)
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        anonymized_names = zf.namelist()
    assert not any("/media/" in name.replace("\\", "/") for name in anonymized_names)

    restored = DocumentRestorer().restore(job.output_path, job.restore_package_path, tmp_path / "restored-image.docx")

    restored_text = "\n".join(chunk.text for chunk in get_handler(restored).extract_text(restored))
    with zipfile.ZipFile(restored, "r") as zf:
        names = zf.namelist()
        document = zf.read("word/document.xml")
        document_rels = zf.read("word/_rels/document.xml.rels")

    assert "Bayer AG" in restored_text
    assert any("/media/" in name.replace("\\", "/") for name in names)
    assert b"<a:blip" in document
    assert b"relationships/image" in document_rels


def test_docx_anonymize_replaces_split_footer_text(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "split-footer.docx"
    doc = Document()
    doc.add_paragraph("Body")
    doc.sections[0].footer.paragraphs[0].text = "Footer"
    doc.save(source)

    footer_xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:p>
    <w:r><w:t>©</w:t></w:r>
    <w:r><w:t xml:space="preserve"> </w:t></w:r>
    <w:r><w:t>msg</w:t></w:r>
    <w:r><w:t xml:space="preserve"> </w:t></w:r>
    <w:r><w:t>s</w:t></w:r>
    <w:r><w:t>ystems</w:t></w:r>
  </w:p>
</w:ftr>"""
    rewritten = tmp_path / "split-footer.rewritten.docx"
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename != "word/footer1.xml":
                zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/footer1.xml", footer_xml)
    shutil.move(rewritten, source)

    job = DocumentJob(source)
    job.findings.append(
        Finding(
            original_text="msg systems",
            replacement_text="Lieferant",
            category="custom",
        )
    )

    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        footer = zf.read("word/footer1.xml").decode("utf-8")

    assert "Lieferant" in footer
    assert "msg" not in footer
    assert "ystems" not in footer


def test_openxml_gallery_detects_and_removes_svg_images(tmp_path: Path) -> None:
    from docx import Document

    from doc_anonymizer.app.formats.common import collect_openxml_images, neutralize_openxml_media

    source = tmp_path / "svg-image.docx"
    target = tmp_path / "svg-image.anonymized.docx"
    doc = Document()
    doc.add_paragraph("Bayer AG")
    doc.save(source)

    with zipfile.ZipFile(source, "r") as zf:
        content_types = zf.read("[Content_Types].xml").decode("utf-8")
        content_types = content_types.replace(
            "</Types>",
            '<Override PartName="/word/media/logo.svg" ContentType="image/svg+xml"/>'
            '<Override PartName="/word/media/photo.wdp" ContentType="image/vnd.ms-photo"/>'
            "</Types>",
        )
    rewritten = tmp_path / "svg-image.rewritten.docx"
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename != "[Content_Types].xml":
                zout.writestr(item, zin.read(item.filename))
        zout.writestr("[Content_Types].xml", content_types)
    shutil.move(rewritten, source)

    with zipfile.ZipFile(source, "a", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "word/header1.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
       xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main">
  <w:p>
    <w:r><w:drawing><a:blip r:embed="rIdSvg"><a:extLst><a:ext><asvg:svgBlip r:embed="rIdSvg"/></a:ext></a:extLst></a:blip></w:drawing></w:r>
    <w:r><w:drawing><a:blip r:embed="rIdWdp"/></w:drawing></w:r>
  </w:p>
</w:hdr>""",
        )
        zf.writestr(
            "word/_rels/header1.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdSvg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/logo.svg"/>
  <Relationship Id="rIdWdp" Type="http://schemas.microsoft.com/office/2007/relationships/hdphoto" Target="media/photo.wdp"/>
</Relationships>""",
        )
        zf.writestr(
            "word/media/logo.svg",
            """<svg xmlns="http://www.w3.org/2000/svg" width="100" height="40"><text x="0" y="20">Logo</text></svg>""",
        )
        zf.writestr("word/media/photo.wdp", b"fake-wdp")

    images = collect_openxml_images(source)
    assert {image.original_file_name for image in images} == {"logo.svg", "photo.wdp"}

    job = DocumentJob(source)
    job.image_replacements = images
    target.write_bytes(source.read_bytes())
    neutralize_openxml_media(source, target, job)

    with zipfile.ZipFile(target, "r") as zf:
        names = zf.namelist()
        header_rels = zf.read("word/_rels/header1.xml.rels")
        header = zf.read("word/header1.xml")
        content_types = zf.read("[Content_Types].xml")

    assert "word/media/logo.svg" not in names
    assert "word/media/photo.wdp" not in names
    assert b"logo.svg" not in content_types
    assert b"photo.wdp" not in content_types
    assert b"relationships/image" not in header_rels
    assert b"relationships/hdphoto" not in header_rels
    assert b"rIdSvg" not in header
    assert b"rIdWdp" not in header
    assert b"<w:drawing" not in header


def test_pptx_anonymize_preserves_run_formatting(tmp_path: Path) -> None:
    from pptx import Presentation

    source = tmp_path / "formatted.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 4000000, 1000000)
    paragraph = textbox.text_frame.paragraphs[0]
    bold_run = paragraph.add_run()
    bold_run.text = "Bayer AG"
    bold_run.font.bold = True
    italic_run = paragraph.add_run()
    italic_run.text = " bleibt sichtbar"
    italic_run.font.italic = True
    prs.save(source)

    job = DocumentJob(source)
    job.findings.append(
        Finding(
            original_text="Bayer AG",
            replacement_text="Firma AG 1",
            category="company",
        )
    )

    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")

    assert "Firma AG 1" in slide_xml
    assert "Bayer AG" not in slide_xml
    assert 'b="1"' in slide_xml
    assert 'i="1"' in slide_xml
    assert " bleibt sichtbar" in slide_xml


def test_pptx_anonymize_replaces_media_payloads_without_removing_relationships(tmp_path: Path) -> None:
    source = tmp_path / "media.pptx"
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="png" ContentType="image/png"/>
</Types>""",
        )
        zf.writestr(
            "ppt/slides/slide1.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:cSld><p:spTree><p:pic><p:blipFill><a:blip r:embed="rIdImage"/></p:blipFill></p:pic></p:spTree></p:cSld>
</p:sld>""",
        )
        zf.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>
</Relationships>""",
        )
        zf.writestr("ppt/media/image1.png", b"not-a-real-image")

    job = DocumentJob(source)
    DocumentAnonymizer().anonymize(job, tmp_path)

    with zipfile.ZipFile(job.output_path, "r") as zf:
        names = zf.namelist()
        rels = zf.read("ppt/slides/_rels/slide1.xml.rels")
        media = zf.read("ppt/media/image1.png")

    assert "ppt/media/image1.png" in names
    assert b"rIdImage" in rels
    assert media != b"not-a-real-image"
    assert media.startswith(b"\x89PNG")
    assert job.image_replacements[0].original_file_name == "image1.png"


def test_pptx_anonymize_removes_shape_image_fill_container(tmp_path: Path) -> None:
    from doc_anonymizer.app.formats.common import neutralize_openxml_media

    source = tmp_path / "shape-fill.pptx"
    target = tmp_path / "shape-fill.anonymized.pptx"
    slide_xml = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:cSld>
    <p:spTree>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="2" name="Picture fill"/></p:nvSpPr>
        <p:spPr>
          <a:blipFill>
            <a:blip r:embed="rIdImage"/>
            <a:stretch><a:fillRect/></a:stretch>
          </a:blipFill>
        </p:spPr>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="3" name="Text"/></p:nvSpPr>
        <p:txBody><a:p><a:r><a:t>Keep me</a:t></a:r></a:p></p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>"""
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="png" ContentType="image/png"/>
</Types>""",
        )
        zf.writestr("ppt/slides/slide1.xml", slide_xml)
        zf.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>
</Relationships>""",
        )
        zf.writestr("ppt/media/image1.png", b"not-a-real-image")

    job = DocumentJob(source)
    target.write_bytes(source.read_bytes())
    neutralize_openxml_media(source, target, job)

    with zipfile.ZipFile(target, "r") as zf:
        anonymized_slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        anonymized_rels = zf.read("ppt/slides/_rels/slide1.xml.rels").decode("utf-8")
        names = zf.namelist()

    assert "ppt/media/image1.png" not in names
    assert "rIdImage" not in anonymized_rels
    assert "Picture fill" not in anonymized_slide
    assert "<a:blipFill" not in anonymized_slide
    assert "Keep me" in anonymized_slide


def test_pptx_anonymize_does_not_replace_technical_guid_attributes(tmp_path: Path) -> None:
    from doc_anonymizer.app.formats.common import replace_in_openxml_text_nodes

    package = tmp_path / "technical-id.pptx"
    slide_xml = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main">
  <p:cSld>
    <p:spTree>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="2" name="Visible 4727">
            <a:extLst>
              <a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}">
                <a16:creationId id="{56803C33-DFCA-4727-8FAE-79FBA6BBAD89}"/>
              </a:ext>
            </a:extLst>
          </p:cNvPr>
        </p:nvSpPr>
        <p:txBody><a:p><a:r><a:t>Visible 4727</a:t></a:r></a:p></p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>"""
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
</Types>""",
        )
        zf.writestr("ppt/slides/slide1.xml", slide_xml)

    replace_in_openxml_text_nodes(
        package,
        [
            Finding(
                original_text="4727",
                replacement_text="Telefon 657",
                category="phone",
            )
        ],
    )

    with zipfile.ZipFile(package, "r") as zf:
        anonymized_slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")

    assert "Visible Telefon 657" in anonymized_slide
    assert '{56803C33-DFCA-4727-8FAE-79FBA6BBAD89}' in anonymized_slide
    assert "{56803C33-DFCA-Telefon 657-8FAE-79FBA6BBAD89}" not in anonymized_slide


def test_pptx_scan_and_anonymize_ignore_technical_uri_attributes(tmp_path: Path) -> None:
    from doc_anonymizer.app.formats.common import extract_openxml_text, replace_in_openxml_text_nodes

    package = tmp_path / "technical-uri.pptx"
    slide_xml = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:graphicFrame>
        <a:graphic>
          <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">
            <a:t>http://example.com/customer</a:t>
          </a:graphicData>
        </a:graphic>
      </p:graphicFrame>
    </p:spTree>
  </p:cSld>
</p:sld>"""
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
</Types>""",
        )
        zf.writestr("ppt/slides/slide1.xml", slide_xml)

    extracted = "\n".join(text for text, _part in extract_openxml_text(package))
    assert "http://example.com/customer" in extracted
    assert "http://schemas.openxmlformats.org/drawingml/2006/diagram" not in extracted

    replace_in_openxml_text_nodes(
        package,
        [
            Finding(
                original_text="http://schemas.openxmlformats.org/drawingml/2006/diagram",
                replacement_text="URL 1",
                category="url",
            ),
            Finding(
                original_text="http://example.com/customer",
                replacement_text="URL 2",
                category="url",
            ),
        ],
    )

    with zipfile.ZipFile(package, "r") as zf:
        anonymized_slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")

    assert 'uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"' in anonymized_slide
    assert "URL 2" in anonymized_slide
    assert 'uri="URL 1"' not in anonymized_slide


def test_openxml_property_name_attributes_are_not_replaced(tmp_path: Path) -> None:
    from doc_anonymizer.app.formats.common import extract_openxml_text, replace_in_openxml_text_nodes

    package = tmp_path / "custom-props.docx"
    custom_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
            xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MediaPersonImageTags">
    <vt:lpwstr>MediaPersonImageTags</vt:lpwstr>
  </property>
</Properties>"""
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
</Types>""",
        )
        zf.writestr("docProps/custom.xml", custom_xml)

    extracted = "\n".join(text for text, _part in extract_openxml_text(package))
    assert "MediaPersonImageTags" in extracted

    replace_in_openxml_text_nodes(
        package,
        [
            Finding(
                original_text="Person",
                replacement_text="Person 1",
                category="person",
            )
        ],
    )

    with zipfile.ZipFile(package, "r") as zf:
        anonymized_props = zf.read("docProps/custom.xml").decode("utf-8")

    assert 'name="MediaPersonImageTags"' in anonymized_props
    assert "MediaPerson 1ImageTags" in anonymized_props


def test_pdf_scan_and_anonymize_with_pdfium(tmp_path: Path) -> None:
    source = tmp_path / "sample.pdf"
    _write_minimal_text_pdf(source, "Herr Max Mustermann")
    job = DocumentJob(source)

    DocumentScanner().scan(job)
    assert any(finding.original_text == "Herr Max Mustermann" for finding in job.findings)

    DocumentAnonymizer().anonymize(job, tmp_path)

    assert job.output_path.exists()
    assert b"Herr Max Mustermann" not in job.output_path.read_bytes()
