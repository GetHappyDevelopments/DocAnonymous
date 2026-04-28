from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from doc_anonymizer.app.formats import FormatHandler, TextChunk
from doc_anonymizer.app.formats.common import (
    extract_openxml_text,
    neutralize_openxml_media,
    replace_in_openxml_text_nodes,
    replace_text,
)


class XlsxHandler(FormatHandler):
    def extract_text(self, path: Path) -> list[TextChunk]:
        wb = load_workbook(path, data_only=False, read_only=False)
        chunks: list[TextChunk] = []
        for ws in wb.worksheets:
            chunks.append(TextChunk(ws.title, f"sheet-name:{ws.title}"))
            if ws.oddHeader.center.text:
                chunks.append(TextChunk(ws.oddHeader.center.text, f"header:{ws.title}"))
            if ws.oddFooter.center.text:
                chunks.append(TextChunk(ws.oddFooter.center.text, f"footer:{ws.title}"))
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value:
                        chunks.append(TextChunk(cell.value, f"sheet:{ws.title}:{cell.coordinate}"))
                    if cell.comment and cell.comment.text:
                        chunks.append(TextChunk(cell.comment.text, f"comment:{ws.title}:{cell.coordinate}"))
        for text, part in extract_openxml_text(path):
            chunks.append(TextChunk(text, part))
        return chunks

    def anonymize(self, source: Path, target: Path, findings: list, job) -> dict[str, Path]:
        wb = load_workbook(source)
        for ws in wb.worksheets:
            ws.title = replace_text(ws.title, findings)[:31]
            for header_footer in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter, ws.firstHeader, ws.firstFooter):
                for side in (header_footer.left, header_footer.center, header_footer.right):
                    if side.text:
                        side.text = replace_text(side.text, findings)
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value:
                        cell.value = replace_text(cell.value, findings)
                    if cell.comment and cell.comment.text:
                        cell.comment.text = replace_text(cell.comment.text, findings)
        wb.save(target)
        replace_in_openxml_text_nodes(target, findings)
        return neutralize_openxml_media(source, target, job)
