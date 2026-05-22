from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

from doc_anonymizer.app.formats import FormatHandler, TextChunk
from doc_anonymizer.app.formats.common import (
    extract_openxml_text,
    replace_openxml_media_with_placeholders,
    replace_in_openxml_text_nodes,
)


class PptxHandler(FormatHandler):
    def extract_text(self, path: Path) -> list[TextChunk]:
        prs = Presentation(str(path))
        chunks: list[TextChunk] = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            for shape_idx, shape in enumerate(slide.shapes):
                if hasattr(shape, "text") and shape.text:
                    chunks.append(TextChunk(shape.text, f"slide:{slide_idx}:shape:{shape_idx}"))
        for text, part in extract_openxml_text(path):
            chunks.append(TextChunk(text, part))
        return chunks

    def anonymize(self, source: Path, target: Path, findings: list, job) -> dict[str, Path]:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        image_files = replace_openxml_media_with_placeholders(source, target, job)
        replace_in_openxml_text_nodes(target, findings)
        return image_files
