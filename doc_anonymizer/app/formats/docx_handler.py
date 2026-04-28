from __future__ import annotations

from pathlib import Path

from docx import Document

from doc_anonymizer.app.formats import FormatHandler, TextChunk
from doc_anonymizer.app.formats.common import (
    extract_openxml_text,
    neutralize_openxml_media,
    replace_in_openxml_text_nodes,
    replace_text,
)


class DocxHandler(FormatHandler):
    def extract_text(self, path: Path) -> list[TextChunk]:
        doc = Document(str(path))
        chunks: list[TextChunk] = []
        for idx, paragraph in enumerate(doc.paragraphs):
            if paragraph.text:
                chunks.append(TextChunk(paragraph.text, f"paragraph:{idx}"))
        for table_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    if cell.text:
                        chunks.append(TextChunk(cell.text, f"table:{table_idx}:{row_idx}:{cell_idx}"))
        for sec_idx, section in enumerate(doc.sections):
            for label, container in (("header", section.header), ("footer", section.footer)):
                for p_idx, paragraph in enumerate(container.paragraphs):
                    if paragraph.text:
                        chunks.append(TextChunk(paragraph.text, f"{label}:{sec_idx}:{p_idx}"))
        seen = {chunk.part for chunk in chunks}
        for text, part in extract_openxml_text(path):
            if part not in seen:
                chunks.append(TextChunk(text, part))
        return chunks

    def anonymize(self, source: Path, target: Path, findings: list, job) -> dict[str, Path]:
        doc = Document(str(source))
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph, findings)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, findings)
        for section in doc.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    self._replace_in_paragraph(paragraph, findings)
        doc.save(str(target))
        replace_in_openxml_text_nodes(target, findings)
        return neutralize_openxml_media(source, target, job)

    @staticmethod
    def _replace_in_paragraph(paragraph, findings: list) -> None:
        if not paragraph.text:
            return
        new_text = replace_text(paragraph.text, findings)
        if new_text == paragraph.text:
            return
        if _replace_within_runs(paragraph, findings):
            return
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)


def _replace_within_runs(paragraph, findings: list) -> bool:
    changed = False
    for finding in findings:
        touched = False
        for run in paragraph.runs:
            if finding.original_text in run.text:
                run.text = run.text.replace(finding.original_text, finding.replacement_text)
                changed = True
                touched = True
        if not touched and finding.original_text in paragraph.text:
            return False
    return changed
